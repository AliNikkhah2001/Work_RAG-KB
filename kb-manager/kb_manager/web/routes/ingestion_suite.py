"""Full ingestion suite: upload/edit/preview/config/save-zip/ingest.

Persian-safe (raw unicode, ensure_ascii=False everywhere in file writes).
Read-only previews; xlsx/docx editable; pdf editing NOT supported.
"""

from __future__ import annotations

import asyncio
import json
import shutil
import uuid
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal

from fastapi import APIRouter, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse

from kb_manager.config import PROJECT_ROOT
from kb_manager.parsers.xlsx_parser import XlsxParser
from kb_manager.chunker.semantic import SemanticChunker

router = APIRouter()

BASE_DIR = PROJECT_ROOT / "data" / "uploads" / "ingestion"
BASE_DIR.mkdir(parents=True, exist_ok=True)
KB_ZIPS_DIR = PROJECT_ROOT / "data" / "kb_zips"
KB_ZIPS_DIR.mkdir(parents=True, exist_ok=True)

MAX_ZIP_BYTES = 500 * 1024 * 1024

_JOBS: dict[str, dict[str, Any]] = {}
_TASKS: dict[str, asyncio.Task] = {}


# ---------------------------------------------------------------------------
# Session helpers
# ---------------------------------------------------------------------------

def _check_sid(sid: str) -> str:
    sid = (sid or "").strip()
    if not sid or not all(c.isalnum() or c in ("-", "_") for c in sid):
        raise HTTPException(status_code=400, detail="Invalid session id")
    return sid


def _session_dir(sid: str) -> Path:
    return BASE_DIR / _check_sid(sid)


def _source_dir(sid: str) -> Path:
    return _session_dir(sid) / "source"


def _session_json_path(sid: str) -> Path:
    return _session_dir(sid) / "session.json"


def _new_session_cfg(sid: str, name: str) -> dict:
    return {
        "id": sid,
        "name": name,
        "created_at": datetime.now(UTC).isoformat(),
        "tags": {},
        "include": {},
        "columns": {},
        "sheet_cache": {},
        "review": {},
    }


def _load_session(sid: str) -> dict:
    p = _session_json_path(sid)
    if not p.exists():
        raise HTTPException(status_code=404, detail="Session not found")
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        raise HTTPException(status_code=500, detail="Error reading session")


def _save_session(sid: str, cfg: dict) -> None:
    _session_dir(sid).mkdir(parents=True, exist_ok=True)
    _session_json_path(sid).write_text(
        json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _safe_resolve(base: Path, rel: str) -> Path:
    """Resolve rel inside base; raise 400 on traversal. Empty rel -> base."""
    rel = (rel or "").strip().lstrip("/")
    # normalize backslashes from clients
    rel = rel.replace("\\", "/")
    if not rel or rel == ".":
        return base
    candidate = (base / rel).resolve()
    base_r = base.resolve()
    if candidate != base_r and base_r not in candidate.parents:
        raise HTTPException(status_code=400, detail="Invalid path")
    return candidate


def _rel_posix(root: Path, p: Path) -> str:
    try:
        return p.resolve().relative_to(root.resolve()).as_posix()
    except Exception:
        return p.name


# ---------------------------------------------------------------------------
# Column-filter helpers (used by ingest + unit tests)
# ---------------------------------------------------------------------------

def get_sheet_config(session_cfg: dict, file_rel: str, sheet_name: str) -> dict:
    """Return {schema_override, target_columns} for a file/sheet ({} if none)."""
    try:
        return (session_cfg.get("columns") or {}).get(file_rel, {}).get(sheet_name, {}) or {}
    except Exception:
        return {}


def filter_sheet_data(
    sheet_data: dict,
    target_columns: list[str] | None = None,
    schema_override: str | None = None,
) -> dict:
    """Return a filtered copy of sheet_data honoring target_columns/schema_override.

    - target_columns: keep only these headers (exact match), in given order.
    - schema_override: replace sheet schema string.
    """
    out = {**sheet_data}
    if target_columns:
        idx_by_h = {h: i for i, h in enumerate(sheet_data.get("headers", []))}
        new_headers = [h for h in target_columns if h in idx_by_h]
        if new_headers:
            new_rows = [
                [r[idx_by_h[h]] if idx_by_h[h] < len(r) else "" for h in new_headers]
                for r in sheet_data.get("rows", [])
            ]
            out = {**out, "headers": new_headers, "rows": new_rows}
    if schema_override:
        out = {**out, "schema": schema_override}
    return out


def _is_included(rel: str, include_map: dict) -> bool:
    """Missing entry => included. Entry False => excluded. Excluded ancestor => excluded."""
    if not include_map:
        return True
    if include_map.get(rel) is False:
        return False
    # ancestor check for recursive excludes
    parts = rel.split("/")
    for i in range(1, len(parts)):
        ancestor = "/".join(parts[:i])
        if include_map.get(ancestor) is False:
            return False
    return True


class SessionFilteredXlsxParser:
    """Wrapper honoring session column config (file_rel -> sheet -> config).

    Wraps XlsxParser.parse: parses normally, then filters sheets listed in
    session_cfg["columns"]. Exposes the same can_parse/parse interface so it
    can be passed as orchestrator parsers={".xlsx": ...}.
    """

    def __init__(self, session_cfg: dict, source_root: str | Path) -> None:
        from kb_manager.parsers.xlsx_parser import XlsxParser

        self._inner = XlsxParser()
        self._cfg = session_cfg or {}
        self._root = Path(source_root)

    def can_parse(self, file_path: str) -> bool:
        return self._inner.can_parse(file_path)

    def parse(self, file_path: str):
        parsed = self._inner.parse(file_path)
        try:
            rel = Path(file_path).resolve().relative_to(self._root.resolve()).as_posix()
        except Exception:
            return parsed
        file_cfg = (self._cfg.get("columns") or {}).get(rel, {})
        if not file_cfg:
            return parsed
        new_sheets = []
        for sh in parsed.sheets or []:
            cfg = file_cfg.get(sh.get("name", ""), {})
            if cfg:
                sh = filter_sheet_data(
                    sh, cfg.get("target_columns"), cfg.get("schema_override")
                )
            new_sheets.append(sh)
        parsed.sheets = new_sheets
        parsed.content = "\n\n".join(self._inner._sheet_to_text(s) for s in new_sheets)
        return parsed


# ---------------------------------------------------------------------------
# Tree
# ---------------------------------------------------------------------------

def _build_node(abs_path: Path, root: Path, cfg: dict) -> dict:
    rel = "" if abs_path.resolve() == root.resolve() else _rel_posix(root, abs_path)
    include_map = cfg.get("include") or {}
    tags_map = cfg.get("tags") or {}
    if abs_path.is_dir():
        children = []
        try:
            entries = sorted(abs_path.iterdir(), key=lambda p: (p.is_file(), p.name))
        except Exception:
            entries = []
        for child in entries:
            if child.name == "__pycache__":
                continue
            children.append(_build_node(child, root, cfg))
        return {
            "path": rel,
            "name": abs_path.name if rel else (cfg.get("name") or root.name),
            "type": "dir",
            "children": children,
            "size": 0,
            "include": include_map.get(rel, True) if rel else True,
            "tags": tags_map.get(rel, []),
        }
    else:
        try:
            size = abs_path.stat().st_size
        except Exception:
            size = 0
        return {
            "path": rel,
            "name": abs_path.name,
            "type": "file",
            "children": [],
            "size": size,
            "include": include_map.get(rel, True),
            "tags": tags_map.get(rel, []),
        }


def _shift_prefix_keys(cfg: dict, old: str, new: str | None, delete: bool = False) -> None:
    """Rename/delete prefixed keys in include/tags/review/columns after fs rename/delete."""
    for key in ("include", "tags", "review"):
        mapping = cfg.get(key) or {}
        updated = {}
        for k, v in mapping.items():
            if k == old or k.startswith(old + "/"):
                if delete:
                    continue
                assert new is not None
                nk = new + k[len(old):]
                updated[nk] = v
            else:
                updated[k] = v
        cfg[key] = updated
    cols = cfg.get("columns") or {}
    updated_c: dict = {}
    for k, v in cols.items():
        if k == old or k.startswith(old + "/"):
            if delete:
                continue
            assert new is not None
            updated_c[new + k[len(old):]] = v
        else:
            updated_c[k] = v
    cfg["columns"] = updated_c


# ---------------------------------------------------------------------------
# 1. Sessions
# ---------------------------------------------------------------------------

@router.post("/sessions")
async def create_session(file: UploadFile = File(...)):
    name = (file.filename or "upload").strip() or "upload"
    if not name.lower().endswith(".zip"):
        raise HTTPException(status_code=400, detail="Please upload a .zip file")
    content = await file.read()
    if len(content) > MAX_ZIP_BYTES:
        raise HTTPException(status_code=400, detail="File exceeds 500MB")
    import io as _io

    try:
        with zipfile.ZipFile(_io.BytesIO(content)) as z:
            bad = z.testzip()
            if bad is not None:
                raise zipfile.BadZipFile(f"corrupt member: {bad}")
            members = z.infolist()
    except zipfile.BadZipFile as e:
        raise HTTPException(status_code=400, detail=f"فایل zip معتبر نیست: {e}")

    sid = uuid.uuid4().hex[:12]
    src = _source_dir(sid)
    src.mkdir(parents=True, exist_ok=True)
    # Extract with zip-slip guard, Persian filenames preserved (utf-8)
    with zipfile.ZipFile(_io.BytesIO(content)) as z:
        for m in members:
            target = _safe_resolve(src, m.filename)
            if m.is_dir():
                target.mkdir(parents=True, exist_ok=True)
            else:
                target.parent.mkdir(parents=True, exist_ok=True)
                with z.open(m) as fsrc, open(target, "wb") as fdst:
                    shutil.copyfileobj(fsrc, fdst)

    stem = Path(name).stem or "session"
    cfg = _new_session_cfg(sid, stem)
    _save_session(sid, cfg)
    return JSONResponse({"session_id": sid, "name": stem})


def _copy_local_tree_into_session(src_path: Path, sid: str) -> int:
    """Copy a server-local dir (or single file) into the session source. Returns file count."""
    dest = _source_dir(sid)
    dest.mkdir(parents=True, exist_ok=True)
    count = 0
    if src_path.is_dir():
        for p in sorted(src_path.rglob("*")):
            if p.is_file():
                rel = p.relative_to(src_path).as_posix()
                target = _safe_resolve(dest, rel)
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(str(p), str(target))
                count += 1
    elif src_path.is_file():
        if src_path.suffix.lower() == ".zip":
            import io as _io

            with zipfile.ZipFile(str(src_path)) as z:
                for m in z.infolist():
                    target = _safe_resolve(dest, m.filename)
                    if m.is_dir():
                        target.mkdir(parents=True, exist_ok=True)
                    else:
                        target.parent.mkdir(parents=True, exist_ok=True)
                        with z.open(m) as fsrc, open(target, "wb") as fdst:
                            shutil.copyfileobj(fsrc, fdst)
                            count += 1
        else:
            target = _safe_resolve(dest, src_path.name)
            shutil.copy2(str(src_path), str(target))
            count = 1
    return count


@router.post("/sessions/from-path")
async def create_session_from_path(payload: dict):
    """Create a session from a server-local path (dir, file, or .zip) — no upload.

    Body: {"path": "D:/data/my_kb" | "./kb-source/KB_9.7.2026", "name": optional}
    """
    raw = str((payload or {}).get("path") or "").strip()
    if not raw:
        raise HTTPException(status_code=400, detail="مسیر لازم است / path required")
    src = Path(raw).expanduser()
    if not src.is_absolute():
        src = (PROJECT_ROOT / src).resolve()
    if not src.exists():
        raise HTTPException(status_code=404, detail=f"مسیر یافت نشد: {raw}")
    sid = uuid.uuid4().hex[:12]
    try:
        count = _copy_local_tree_into_session(src, sid)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"خطا در خواندن مسیر: {e}")
    if count == 0:
        raise HTTPException(status_code=400, detail="Empty path")
    name = str((payload or {}).get("name") or src.stem or "local").strip() or "local"
    cfg = _new_session_cfg(sid, Path(name).stem)
    cfg["source_path"] = str(src)
    _save_session(sid, cfg)
    return JSONResponse({"session_id": sid, "name": cfg["name"], "files": count, "source_path": str(src)})


@router.post("/sessions/{sid}/export-path")
async def export_session_to_path(sid: str, payload: dict):
    """Write session source tree to a server-local directory (no download).

    Body: {"path": "D:/data/out_kb"} — created if missing, files overwritten.
    """
    _load_session(sid)
    raw = str((payload or {}).get("path") or "").strip()
    if not raw:
        raise HTTPException(status_code=400, detail="مسیر لازم است / path required")
    dest = Path(raw).expanduser()
    if not dest.is_absolute():
        dest = (PROJECT_ROOT / dest).resolve()
    root = _source_dir(sid)
    try:
        dest.mkdir(parents=True, exist_ok=True)
        count = 0
        for p in sorted(root.rglob("*")):
            if p.is_file():
                rel = p.relative_to(root).as_posix()
                # guard against .. in stored names
                target = dest / Path(rel)
                if dest.resolve() not in target.resolve().parents and target.resolve() != dest.resolve():
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(str(p), str(target))
                count += 1
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"خطا در نوشتن مسیر: {e}")
    return JSONResponse({"ok": True, "path": str(dest), "files": count})


@router.get("/sessions")
async def list_sessions():
    out = []
    if BASE_DIR.exists():
        for d in sorted(BASE_DIR.iterdir()):
            if not d.is_dir():
                continue
            p = d / "session.json"
            if not p.exists():
                continue
            try:
                cfg = json.loads(p.read_text(encoding="utf-8"))
                src = d / "source"
                n_files = sum(1 for _ in src.rglob("*") if _.is_file()) if src.exists() else 0
                out.append({
                    "id": cfg.get("id", d.name),
                    "name": cfg.get("name", d.name),
                    "created_at": cfg.get("created_at", ""),
                    "file_count": n_files,
                })
            except Exception:
                continue
    return JSONResponse(out)


# ---------------------------------------------------------------------------
# 2. Tree
# ---------------------------------------------------------------------------

@router.get("/sessions/{sid}/tree")
async def get_tree(sid: str):
    cfg = _load_session(sid)
    root = _source_dir(sid)
    root.mkdir(parents=True, exist_ok=True)
    return JSONResponse(_build_node(root, root, cfg))


@router.post("/sessions/{sid}/tree/{op}")
async def tree_op(sid: str, op: Literal["rename", "delete", "mkdir", "tag", "include"], payload: dict):
    cfg = _load_session(sid)
    root = _source_dir(sid)
    rel = str(payload.get("path") or "").replace("\\", "/").strip().lstrip("/")

    if op == "mkdir":
        if not rel:
            raise HTTPException(status_code=400, detail="New folder path is required")
        target = _safe_resolve(root, rel)
        target.mkdir(parents=True, exist_ok=True)
        return JSONResponse({"ok": True, "path": rel})

    if op == "rename":
        new_name = str(payload.get("new_name") or "").strip().replace("/", "").replace("\\", "")
        if not rel or not new_name:
            raise HTTPException(status_code=400, detail="Path and new name are required")
        target = _safe_resolve(root, rel)
        if not target.exists():
            raise HTTPException(status_code=404, detail="Path not found")
        dest = target.parent / new_name
        if dest.exists():
            raise HTTPException(status_code=400, detail="A file with this name already exists")
        target.rename(dest)
        new_rel = _rel_posix(root, dest)
        _shift_prefix_keys(cfg, rel, new_rel)
        _save_session(sid, cfg)
        return JSONResponse({"ok": True, "path": new_rel})

    if op == "delete":
        if not rel:
            raise HTTPException(status_code=400, detail="Path is required")
        target = _safe_resolve(root, rel)
        if not target.exists():
            raise HTTPException(status_code=404, detail="Path not found")
        if target.is_dir() and target.resolve() != root.resolve():
            shutil.rmtree(target)
        elif target.is_file():
            target.unlink()
        else:
            raise HTTPException(status_code=400, detail="Cannot delete session root")
        _shift_prefix_keys(cfg, rel, None, delete=True)
        _save_session(sid, cfg)
        return JSONResponse({"ok": True})

    if op == "tag":
        tags = payload.get("tags", [])
        cfg.setdefault("tags", {})[rel] = tags
        _save_session(sid, cfg)
        return JSONResponse({"ok": True, "path": rel, "tags": tags})

    if op == "include":
        include_val = payload.get("include", True)
        recursive = bool(payload.get("recursive", False))
        cfg.setdefault("include", {})[rel] = bool(include_val)
        if recursive:
            target = _safe_resolve(root, rel)
            if target.is_dir():
                for p in target.rglob("*"):
                    r = _rel_posix(root, p)
                    cfg["include"][r] = bool(include_val)
        _save_session(sid, cfg)
        return JSONResponse({"ok": True, "path": rel, "include": bool(include_val)})

    raise HTTPException(status_code=400, detail="Invalid operation")


# ---------------------------------------------------------------------------
# 3. File preview (read-only)
# ---------------------------------------------------------------------------

@router.get("/sessions/{sid}/file")
async def preview_file(sid: str, path: str = Query(...)):
    _load_session(sid)
    root = _source_dir(sid)
    rel = path.replace("\\", "/").strip().lstrip("/")
    target = _safe_resolve(root, rel)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    suffix = target.suffix.lower()
    # Extensionless/misnamed workbooks (e.g. "Name(V.3)"): sniff ZIP magic
    try:
        with open(str(target), "rb") as _f:
            _magic = _f.read(4)
    except Exception:
        _magic = b""
    is_xlsx = suffix == ".xlsx" or _magic == b"PK\x03\x04"

    if is_xlsx:
        from kb_manager.parsers.xlsx_parser import (
            _detect_schema,
            _normalize_col,
        )
        from kb_manager.parsers.xlsx_parser import XlsxParser

        try:
            parsed = XlsxParser().parse(str(target))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error reading Excel: {e}")
        # schema column sets for overlap
        try:
            from kb_manager.parsers.xlsx_parser import (
                SCHEMA_A_COLUMNS,
                SCHEMA_B_COLUMNS,
                SCHEMA_C_COLUMNS,
            )
            schema_sets = {
                "reason_codes": {_normalize_col(c) for c in SCHEMA_A_COLUMNS},
                "crm_qa": {_normalize_col(c) for c in SCHEMA_B_COLUMNS} | {"question", "model", "briefanswer", "answer", "keyword"},
                "articles": {_normalize_col(c) for c in SCHEMA_C_COLUMNS},
            }
        except Exception:
            schema_sets = {}
        sheets = []
        for sh in parsed.sheets or []:
            guess = sh.get("schema") or _detect_schema(sh.get("headers", []))
            norm = {_normalize_col(h) for h in sh.get("headers", []) if h}
            overlap = len(norm & schema_sets.get(guess, set())) if guess in schema_sets else 0
            sheets.append({
                "name": sh.get("name"),
                "headers": sh.get("headers"),
                "rows_preview": (sh.get("rows") or [])[:20],
                "row_count": len(sh.get("rows") or []),
                "schema_guess": guess,
                "overlap": overlap,
            })
        return JSONResponse({"type": "xlsx", "path": rel, "sheets": sheets})

    if suffix == ".docx":
        from kb_manager.parsers.docx_parser import DocxParser

        try:
            parsed = DocxParser().parse(str(target))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"خطا در خواندن docx: {e}")
        return JSONResponse({
            "type": "docx", "path": rel,
            "title": parsed.title,
            "sections": parsed.sections or [],
            "text_preview": (parsed.content or "")[:4000],
        })

    if suffix == ".pdf":
        from kb_manager.parsers.pdf_parser import PdfParser

        try:
            parsed = PdfParser().parse(str(target))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"خطا در خواندن pdf: {e}")
        pages = [p for p in (parsed.content or "").split("\n\n") if p.strip()][:3]
        return JSONResponse({
            "type": "pdf", "path": rel,
            "page_count": (parsed.metadata or {}).get("page_count"),
            "pages_preview": pages,
            "sections_preview": (parsed.sections or [])[:5],
        })

    # fallback: text preview
    try:
        text = target.read_text(encoding="utf-8", errors="replace")
        return JSONResponse({"type": "text", "path": rel, "text_preview": text[:4000]})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"پیش‌نمایش پشتیبانی نمی‌شود: {e}")


# ---------------------------------------------------------------------------
# 3b. Chunking visualization (how table rows map to chunks)
# ---------------------------------------------------------------------------

@router.get("/sessions/{sid}/file/chunking")
async def chunking_visualization(sid: str, path: str = Query(...)):
    """Return detailed chunking mapping for an XLSX file.
    
    Shows which columns/rows become which chunks, with color coding per chunk.
    """
    _load_session(sid)
    root = _source_dir(sid)
    rel = path.replace("\\", "/").strip().lstrip("/")
    target = _safe_resolve(root, rel)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    # Accept .xlsx plus extensionless/misnamed workbooks (e.g. "Name(V.3)"): sniff ZIP magic
    try:
        with open(str(target), "rb") as _f:
            _magic = _f.read(4)
    except Exception:
        _magic = b""
    if target.suffix.lower() != ".xlsx" and _magic != b"PK\x03\x04":
        raise HTTPException(status_code=400, detail="Chunking visualization only for .xlsx files")

    try:
        parsed = XlsxParser().parse(str(target))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error parsing xlsx: {e}")

    # Determine doc_type for each sheet (same logic as orchestrator)
    from kb_manager.parsers.xlsx_parser import _detect_schema
    sheet_doc_types = {}
    for sh in parsed.sheets or []:
        schema = sh.get("schema") or _detect_schema(sh.get("headers", []))
        # Map schema to doc_type (mirrors pipeline/orchestrator.py)
        if schema == "crm_qa":
            doc_type = "qa_pair"
        elif schema == "reason_codes":
            doc_type = "reason_detail"
        elif schema == "articles":
            doc_type = "article"
        elif schema == "glossary":
            doc_type = "glossary"
        elif schema == "loan_catalog":
            doc_type = "loan_catalog"
        elif schema == "staff_profile":
            doc_type = "staff_profile"
        elif schema == "timeline":
            doc_type = "timeline"
        elif schema == "single_col_list":
            doc_type = "single_col_list"
        else:
            doc_type = "body"
        sheet_doc_types[sh.get("name", "")] = doc_type

    # Simulate chunking for each sheet
    chunker = SemanticChunker(
        parent_scope="sheet",
        parent_max_tokens=1536,
        dedup_questions=True,
    )
    # Create mock metadata as orchestrator would
    mock_metadata = {
        "doc_type": "qa_pair",  # default, overridden per sheet
        "parent_scope": "sheet",
        "sheets": [],
    }

    # Color palette for chunks (each row = one chunk for tabular types)
    CHUNK_COLORS = [
        "#E8F5E9", "#E3F2FD", "#FFF3E0", "#FCE4EC", "#F3E5F5", "#E0F2F1",
        "#FFF8E1", "#F1F8E9", "#EDE7F6", "#E0E0E0", "#FBE9E7", "#E8EAFC",
    ]

    visualization = {
        "file_path": rel,
        "sheets": [],
        "color_legend": [],
    }

    chunk_index = 0
    for sh in parsed.sheets or []:
        sheet_name = sh.get("name", "")
        headers = sh.get("headers", [])
        rows = sh.get("rows", [])
        schema = sh.get("schema") or _detect_schema(headers)
        doc_type = sheet_doc_types.get(sheet_name, "body")

        # Only tabular types get row-wise chunks (mirrors SemanticChunker.chunk)
        is_tabular = doc_type in ("qa_pair", "reason_detail", "article", "single_col_list", "glossary", "loan_catalog", "staff_profile", "timeline")

        # Header -> chunk label mapping (labels used inside chunk content)
        fa_names = SemanticChunker._FIELD_NAMES_FA
        if schema == "reason_code":
            _label_order = [
                ("reason_code", "کد دلیل"),
                ("model", "مدل"),
                ("brief_explanation", "توضیح کوتاه"),
                ("detailed_explanation", "توضیح کامل"),
            ]
        else:
            _label_order = [
                ("question", "سوال"),
                ("پرسش", "سوال"),
                ("سوال", "سوال"),
                ("متن سوال", "سوال"),
                ("متن_سوال", "سوال"),
                ("briefanswer", "پاسخ کوتاه"),
                ("brief_answer", "پاسخ کوتاه"),
                ("answer", "پاسخ کامل"),
                ("پاسخ", "پاسخ کامل"),
                ("متن پاسخ", "پاسخ کامل"),
                ("متن_پاسخ", "پاسخ کامل"),
                ("keyword", "کلیدواژه‌ها"),
                ("keywords", "کلیدواژه‌ها"),
            ]
        _order_labels = {k: v for k, v in _label_order}
        field_mappings = [
            {"header": h, "label": _order_labels.get(h, _order_labels.get(h.lower(), fa_names.get(h.lower(), h)))}
            for h in headers
        ]

        sheet_viz = {
            "name": sheet_name,
            "schema": schema,
            "doc_type": doc_type,
            "is_tabular": is_tabular,
            "headers": headers,
            "field_mappings": field_mappings,
            "total_rows": len(rows),
            "chunks": [],
            "skipped_rows": 0,
        }

        if is_tabular and rows:
            # Simulate chunking for this sheet
            mock_metadata["doc_type"] = doc_type
            mock_metadata["sheets"] = [sh]
            try:
                chunks = chunker.chunk("", mock_metadata)
                # Filter chunks belonging to this sheet
                sheet_chunks = [c for c in chunks if c.metadata.get("sheet_name") == sheet_name]
            except Exception:
                sheet_chunks = []

            # Non-empty row indices, in order (chunker emits one chunk per
            # non-empty row, in order, unless QA dedup drops duplicates)
            nonempty_idx = [
                i for i, r in enumerate(rows)
                if any(v and str(v).strip() for v in (r or []))
            ]
            order_aligned = len(sheet_chunks) == len(nonempty_idx)
            chunk_by_row = {ri: sheet_chunks[k] for k, ri in enumerate(nonempty_idx)} if order_aligned else {}

            for row_idx, row in enumerate(rows):
                # Build field map for this row
                fields = {}
                for header, value in zip(headers, row):
                    if value and str(value).strip():
                        fields[header] = str(value).strip()

                # Skip only when ALL selected columns are empty (partial rows kept)
                skipped = not fields
                skip_reason = "All selected columns empty" if skipped else ""

                # Format content as chunker would
                if is_tabular and not skipped:
                    # Find matching chunk: order-aligned first, then key match
                    chunk = chunk_by_row.get(row_idx)
                    if chunk is None:
                        for c in sheet_chunks:
                            cf = c.metadata.get("fields", {})
                            # Match by question or reason_code
                            if doc_type == "qa_pair":
                                cq = cf.get("question", "") or cf.get("پرسش", "") or cf.get("سوال", "")
                                rq = fields.get("question", "") or fields.get("پرسش", "") or fields.get("سوال", "")
                                if cq and rq and cq == rq:
                                    chunk = c
                                    break
                            elif doc_type == "reason_detail":
                                cq = cf.get("reason_code", "")
                                rq = fields.get("reason_code", "")
                                if cq and rq and cq == rq:
                                    chunk = c
                                    break
                    if chunk:
                        content = chunk.content
                    elif doc_type in ("qa_pair", "reason_detail"):
                        # Fallback: format ourselves
                        content = _format_qa_content(fields, schema)
                    else:
                        # Fallback mirrors chunker pipe format for other tabular types
                        content = " | ".join(
                            f"{h}: {fields[h]}" for h in headers if h in fields
                        )
                else:
                    content = " | ".join(f"{h}: {v}" for h, v in zip(headers, row) if v and str(v).strip())

                color = CHUNK_COLORS[chunk_index % len(CHUNK_COLORS)]
                if chunk_index == 0:
                    visualization["color_legend"].append({"color": color, "label": f"Chunk {chunk_index + 1}"})

                sheet_viz["chunks"].append({
                    "row_index": row_idx,
                    "skipped": skipped,
                    "skip_reason": skip_reason,
                    "fields": fields,
                    "content_preview": content[:500] + ("..." if len(content) > 500 else ""),
                    "content_full": content,
                    "color": color if not skipped else "#FFEBEE",
                    "chunk_type": doc_type if not skipped else "skipped",
                    "token_estimate": len(content.split()) * 1.3 if content else 0,
                })
                if not skipped:
                    chunk_index += 1

            sheet_viz["skipped_rows"] = sum(1 for c in sheet_viz["chunks"] if c["skipped"])

        elif rows:
            # Non-tabular (body) sheets: rows merge into flowing text and are
            # split structurally (articles/sections/pipe-rows/newlines). Show
            # the resulting chunks and map each row to the chunk containing it.
            try:
                sheet_text = XlsxParser()._sheet_to_text({
                    "name": sheet_name, "headers": headers, "rows": rows, "schema": schema,
                })
            except Exception:
                sheet_text = ""
            try:
                struct_chunks = chunker._chunk_structural(
                    sheet_text,
                    {"document_id": "preview", "doc_type": "body", "sheet_name": sheet_name},
                )
            except Exception:
                struct_chunks = []
            struct_chunks = [c for c in struct_chunks if not c.metadata.get("is_parent")]
            for row_idx, row in enumerate(rows):
                cells = [str(v).strip() for v in (row or []) if v and str(v).strip()]
                sig = next((c for c in cells if len(c) >= 8), (cells[0] if cells else ""))
                sig_key = sig[:40]
                in_chunk = next(
                    (k for k, c in enumerate(struct_chunks) if sig_key and sig_key in c.content),
                    None,
                )
                fields = {h: str(v).strip() for h, v in zip(headers, row) if v and str(v).strip()}
                if in_chunk is None and not fields:
                    skipped, skip_reason, content = True, "All selected columns empty", ""
                elif in_chunk is None:
                    skipped, skip_reason = False, ""
                    content = " | ".join(f"{h}: {fields[h]}" for h in headers if h in fields)
                else:
                    skipped, skip_reason = False, ""
                    content = struct_chunks[in_chunk].content
                color = CHUNK_COLORS[in_chunk % len(CHUNK_COLORS)] if in_chunk is not None else "#E0E0E0"
                sheet_viz["chunks"].append({
                    "row_index": row_idx,
                    "skipped": skipped,
                    "skip_reason": skip_reason,
                    "fields": fields,
                    "content_preview": (content[:500] + ("..." if len(content) > 500 else "")) if content else "",
                    "content_full": content,
                    "color": color if not skipped else "#FFEBEE",
                    "chunk_type": f"body (chunk {in_chunk + 1}/{len(struct_chunks)})" if in_chunk is not None else ("skipped" if skipped else "body"),
                    "token_estimate": len(content.split()) * 1.3 if content else 0,
                })
            sheet_viz["structural_chunks"] = len(struct_chunks)

        visualization["sheets"].append(sheet_viz)

    return JSONResponse(visualization)


# Helper function for formatting (copied from chunker)
def _format_qa_content(fields: dict[str, str], schema: str) -> str:
    """Format a QA/reason-code row with field names."""
    lines = []
    if schema == "reason_code":
        field_order = [
            ("reason_code", "کد دلیل"),
            ("model", "مدل"),
            ("brief_explanation", "توضیح کوتاه"),
            ("detailed_explanation", "توضیح کامل"),
        ]
    else:
        field_order = [
            ("question", "سوال"),
            ("پرسش", "سوال"),
            ("سوال", "سوال"),
            ("متن سوال", "سوال"),
            ("متن_سوال", "سوال"),
            ("briefanswer", "پاسخ کوتاه"),
            ("brief_answer", "پاسخ کوتاه"),
            ("answer", "پاسخ کامل"),
            ("پاسخ", "پاسخ کامل"),
            ("متن پاسخ", "پاسخ کامل"),
            ("متن_پاسخ", "پاسخ کامل"),
            ("keyword", "کلیدواژه‌ها"),
            ("keywords", "کلیدواژه‌ها"),
        ]
    for key, label in field_order:
        if key in fields:
            lines.append(f"{label}: {fields[key]}")
    seen = {k for k, _ in field_order}
    for key, value in fields.items():
        if key not in seen:
            lines.append(f"{key}: {value}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 4. File edit
# ---------------------------------------------------------------------------

@router.put("/sessions/{sid}/file")
async def edit_file(sid: str, path: str = Query(...), payload: dict | None = None):
    _load_session(sid)
    root = _source_dir(sid)
    rel = (path or "").replace("\\", "/").strip().lstrip("/")
    target = _safe_resolve(root, rel)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    payload = payload or {}
    suffix = target.suffix.lower()

    if suffix == ".xlsx":
        sheets_payload = payload.get("sheets") or {}
        # Accept flat single-sheet payload from the editor too:
        # {sheet, headers, rows} -> {sheets: {sheet: {headers, rows}}}
        if (not isinstance(sheets_payload, dict) or not sheets_payload) and isinstance(payload.get("headers"), list):
            flat_sheet = payload.get("sheet")
            if flat_sheet:
                sheets_payload = {str(flat_sheet): {"headers": payload.get("headers", []), "rows": payload.get("rows", [])}}
        if not isinstance(sheets_payload, dict) or not sheets_payload:
            raise HTTPException(status_code=400, detail="Sheet data is required")
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(target))
            try:
                for sheet_name, data in sheets_payload.items():
                    headers = (data or {}).get("headers", [])
                    rows = (data or {}).get("rows", [])
                    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
                    if ws.max_row and ws.max_column:
                        ws.delete_rows(1, ws.max_row)
                    ws.append(list(headers))
                    for r in rows or []:
                        ws.append(list(r))
            finally:
                pass
            wb.save(str(target))
            wb.close()
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error saving xlsx: {e}")
        return JSONResponse({"ok": True, "path": rel, "message": "Excel file saved"})

    if suffix == ".docx":
        sections = payload.get("sections")
        text = payload.get("text")
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument()
            if sections:
                for sec in sections:
                    heading = (sec or {}).get("heading", "")
                    body = (sec or {}).get("text", "")
                    if heading:
                        doc.add_heading(str(heading), level=2)
                    for para in str(body or "").split("\n"):
                        if para.strip():
                            doc.add_paragraph(para)
            elif text is not None:
                for para in str(text).split("\n\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())
            else:
                raise HTTPException(status_code=400, detail="Sections or text are required")
            doc.save(str(target))
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error saving docx: {e}")
        return JSONResponse({"ok": True, "path": rel, "message": "Docx file saved"})

    if suffix == ".pdf":
        raise HTTPException(status_code=400, detail="PDF editing is not supported")

    raise HTTPException(status_code=400, detail="Editing this file type is not supported")


# ---------------------------------------------------------------------------
# 5. Config
# ---------------------------------------------------------------------------

@router.get("/sessions/{sid}/config")
async def get_config(sid: str):
    cfg = _load_session(sid)
    return JSONResponse({
        "include": cfg.get("include", {}),
        "columns": cfg.get("columns", {}),
        "tags": cfg.get("tags", {}),
    })


@router.put("/sessions/{sid}/config")
async def put_config(sid: str, payload: dict):
    cfg = _load_session(sid)
    for key in ("include", "columns", "tags"):
        if key in payload:
            if not isinstance(payload[key], dict):
                raise HTTPException(status_code=400, detail=f"Field {key} must be an object")
            cfg[key] = payload[key]
    _save_session(sid, cfg)
    return JSONResponse({
        "ok": True,
        "include": cfg.get("include", {}),
        "columns": cfg.get("columns", {}),
        "tags": cfg.get("tags", {}),
    })


# ---------------------------------------------------------------------------
# 5b. Review (per-document chunking approval gate)
# ---------------------------------------------------------------------------

def _required_review_files(sid: str, cfg: dict) -> list[str]:
    """All included files that must be reviewed before continuing."""
    root = _source_dir(sid)
    include_map = cfg.get("include") or {}
    required: list[str] = []
    if root.exists():
        for p in sorted(root.rglob("*")):
            if p.is_file():
                rel = _rel_posix(root, p)
                if _is_included(rel, include_map):
                    required.append(rel)
    return required


@router.get("/sessions/{sid}/review")
async def get_review(sid: str):
    cfg = _load_session(sid)
    required = _required_review_files(sid, cfg)
    reviewed = cfg.get("review") or {}
    remaining = [r for r in required if r not in reviewed]
    return JSONResponse({
        "required": required,
        "reviewed": reviewed,
        "remaining": remaining,
        "total": len(required),
        "reviewed_count": len(required) - len(remaining),
        "done": len(remaining) == 0,
    })


@router.post("/sessions/{sid}/review")
async def post_review(sid: str, payload: dict):
    cfg = _load_session(sid)
    rel = str((payload or {}).get("path") or "").replace("\\", "/").strip().lstrip("/")
    if not rel:
        raise HTTPException(status_code=400, detail="Path is required")
    root = _source_dir(sid)
    target = _safe_resolve(root, rel)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    approved = bool((payload or {}).get("approved", True))
    review = cfg.setdefault("review", {})
    if approved:
        review[rel] = {
            "approved_at": datetime.now(UTC).isoformat(),
            "headers": ((cfg.get("columns") or {}).get(rel) or {}),
        }
    else:
        review.pop(rel, None)
    _save_session(sid, cfg)
    required = _required_review_files(sid, cfg)
    remaining = [r for r in required if r not in review]
    return JSONResponse({
        "ok": True,
        "path": rel,
        "approved": approved,
        "remaining": remaining,
        "total": len(required),
        "reviewed_count": len(required) - len(remaining),
        "done": len(remaining) == 0,
    })


# ---------------------------------------------------------------------------
# 6. Save zip
# ---------------------------------------------------------------------------

@router.post("/sessions/{sid}/save-zip")
async def save_zip(sid: str):
    cfg = _load_session(sid)
    root = _source_dir(sid)
    if not root.exists():
        raise HTTPException(status_code=404, detail="Session source not found")
    safe_name = "".join(c if (c.isalnum() or c in ("-", "_", " ")) else "_" for c in (cfg.get("name") or "session")).strip() or "session"
    zip_name = f"{safe_name}-{sid}.zip"
    zip_path = KB_ZIPS_DIR / zip_name
    KB_ZIPS_DIR.mkdir(parents=True, exist_ok=True)
    count = 0
    with zipfile.ZipFile(str(zip_path), "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(root.rglob("*")):
            if p.is_file():
                arc = p.resolve().relative_to(root.resolve()).as_posix()
                z.write(str(p), arc)
                count += 1
    size = zip_path.stat().st_size
    return JSONResponse({"zip_path": str(zip_path), "size": size, "files": count})


# ---------------------------------------------------------------------------
# 7. Ingest (isolated DB, column filtering, best-effort benchmark)
# ---------------------------------------------------------------------------

def _sanitize_db_name(raw: str | None, fallback: str) -> str:
    raw = (raw or "").strip() or f"kb_{fallback}.db"
    base = Path(raw).name
    if not base.lower().endswith(".db"):
        base += ".db"
    safe = "".join(c if (c.isalnum() or c in ("-", "_", ".")) else "_" for c in base)
    return safe or f"kb_{fallback}.db"


async def _run_ingest_job(
    job_id: str,
    sid: str,
    db_path: Path,
    job_type: str,
    parent_scope: str,
    dataset_mode: str,
) -> None:
    job = _JOBS[job_id]
    try:
        from kb_manager.config import DatabaseConfig
        from kb_manager.models.database import Database

        cfg = _load_session(sid)
        root = _source_dir(sid)
        include_map = cfg.get("include") or {}

        iso_db = Database(DatabaseConfig(sqlite_path=str(db_path)))
        await iso_db.create_tables()

        from kb_manager.chunker.semantic import SemanticChunker
        from kb_manager.pipeline.orchestrator import PipelineOrchestrator

        scope = parent_scope if parent_scope in ("sheet", "document") else "sheet"
        chunker = SemanticChunker(parent_scope=scope)
        xlsx = SessionFilteredXlsxParser(cfg, root)
        orch = PipelineOrchestrator(database=iso_db, chunker=chunker, parsers={".xlsx": xlsx})

        # Honor include map by wrapping _scan_files
        _orig_scan = orch._scan_files

        def _filtered_scan(source_dir: str) -> list[str]:
            files = _orig_scan(source_dir)
            out = []
            for f in files:
                try:
                    rel = Path(f).resolve().relative_to(root.resolve()).as_posix()
                except Exception:
                    out.append(f)
                    continue
                if _is_included(rel, include_map):
                    out.append(f)
            return out

        orch._scan_files = _filtered_scan  # type: ignore[method-assign]

        job["status"] = "running"
        if job_type == "incremental":
            summary = await orch.run_incremental(str(root.resolve()))
        else:
            summary = await orch.run_full_rebuild(str(root.resolve()))
        job["summary"] = summary.to_dict()
        try:
            await iso_db.close()
        except Exception:
            pass

        # Best-effort benchmark on data/test_questions.json
        bench: dict[str, Any] = {"attempted": False}
        try:
            tq = PROJECT_ROOT / "data" / "test_questions.json"
            if tq.exists():
                from kb_manager.evaluation.benchmark import BenchmarkRunner

                dataset = json.loads(tq.read_text(encoding="utf-8"))
                runner = BenchmarkRunner(lambda q, k: [], top_k=5, version=f"ingestion-{sid}")
                result = runner.run(dataset if isinstance(dataset, list) else [])
                bench = {"attempted": True, "total": result.total_queries, "status": "done"}
        except Exception as e:
            bench = {"attempted": True, "status": "skipped", "error": str(e)[:300]}
        job["benchmark"] = bench
        job["status"] = "completed"
        job["completed_at"] = datetime.now(UTC).isoformat()
    except Exception as e:
        job["status"] = "failed"
        job["error"] = str(e)[:2000]
        job["completed_at"] = datetime.now(UTC).isoformat()


@router.post("/sessions/{sid}/ingest")
async def start_ingest(sid: str, payload: dict | None = None):
    cfg = _load_session(sid)
    payload = payload or {}
    db_name = _sanitize_db_name(payload.get("db_name"), cfg.get("name") or sid)
    job_type = str(payload.get("job_type") or "full_rebuild")
    if job_type not in ("full_rebuild", "incremental"):
        job_type = "full_rebuild"
    parent_scope = str(payload.get("parent_scope") or "sheet")
    dataset_mode = str(payload.get("dataset_mode") or "default")

    db_path = PROJECT_ROOT / "data" / db_name
    db_path.parent.mkdir(parents=True, exist_ok=True)

    job_id = uuid.uuid4().hex[:12]
    _JOBS[job_id] = {
        "job_id": job_id,
        "session_id": sid,
        "status": "pending",
        "db_path": str(db_path),
        "job_type": job_type,
        "parent_scope": parent_scope,
        "dataset_mode": dataset_mode,
        "started_at": datetime.now(UTC).isoformat(),
        "completed_at": None,
        "error": "",
    }
    task = asyncio.create_task(_run_ingest_job(job_id, sid, db_path, job_type, parent_scope, dataset_mode))
    _TASKS[job_id] = task
    task.add_done_callback(lambda t, jid=job_id: _TASKS.pop(jid, None))
    return JSONResponse({"job_id": job_id, "db_path": str(db_path)}, status_code=202)


@router.get("/jobs/{job_id}")
async def job_status(job_id: str):
    job = _JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return JSONResponse({k: v for k, v in job.items()})


@router.get("", response_class=HTMLResponse)
async def ingestion_page(request: Request):
    """Serve the Persian RTL ingestion suite UI."""
    from kb_manager.web.deps import templates

    return templates.TemplateResponse(request, "ingestion.html", {})
